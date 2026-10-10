import io
from datetime import datetime

import openpyxl
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt

from .evaluation import build_evaluation


def set_chinese_font(run, font_name="宋体", size=11, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold


def generate_docx_report(username, rock_type, items, summary, is_demo=False):
    doc = Document()
    title = doc.add_heading("岩心智能识别与储层物性评价报告", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for run in title.runs:
        set_chinese_font(run, "黑体", 16, True)

    doc.add_paragraph(f"检测用户：{username}")
    doc.add_paragraph(f"岩石类型：{rock_type}")
    doc.add_paragraph(f"检测时间：{datetime.now():%Y-%m-%d %H:%M:%S}")
    doc.add_paragraph(f"样本数量：{summary['count']} 张")
    if is_demo:
        doc.add_paragraph("演示模型结果：仅用于功能演示，不代表真实岩心物性。")

    doc.add_heading("一、检测结果汇总", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "序号"
    headers[1].text = "文件名"
    headers[2].text = "孔隙度(%)"
    headers[3].text = "渗透率(mD)"
    for index, item in enumerate(items, 1):
        row = table.add_row().cells
        row[0].text = str(index)
        row[1].text = item.get("filename") or f"样本{index}"
        row[2].text = f"{float(item['porosity']):.2f}"
        row[3].text = f"{float(item['permeability']):.2f}"

    doc.add_heading("二、统计特征", level=1)
    doc.add_paragraph(f"平均孔隙度：{summary['avg_porosity']}%")
    doc.add_paragraph(f"平均渗透率：{summary['avg_permeability']} mD")
    doc.add_paragraph(f"孔隙度范围：{summary['min_porosity']}% ~ {summary['max_porosity']}%")
    doc.add_paragraph(f"渗透率范围：{summary['min_permeability']} mD ~ {summary['max_permeability']} mD")
    for key, label, unit in (("porosity", "孔隙度", "%"), ("permeability", "渗透率", "mD")):
        stats = summary["metrics"][key]
        cv = f"{stats['cv'] * 100:.2f}%" if stats["cv"] is not None else "不适用"
        doc.add_paragraph(f"{label}中位数：{stats['median']:.2f} {unit}；总体标准差：{stats['std']:.2f} {unit}；变异系数：{cv}")

    evaluation = build_evaluation(summary["avg_porosity"], summary["avg_permeability"])
    doc.add_heading("三、综合评价", level=1)
    doc.add_paragraph(f"孔隙度评价：{evaluation['porosity_level']}，{evaluation['porosity_desc']}")
    doc.add_paragraph(f"渗透率评价：{evaluation['permeability_level']}，{evaluation['permeability_desc']}")
    doc.add_paragraph(f"均值物性参考等级：{evaluation['reservoir_class']}")
    doc.add_heading("四、图表分析总结", level=1)
    for insight in summary["insights"]:
        doc.add_paragraph(insight)
    doc.add_paragraph(summary["analysis_note"])

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def export_history_to_xlsx(username, records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "岩心历史记录"
    ws.append(["序号", "岩石类型", "孔隙度(%)", "渗透率(mD)", "检测时间", "导出用户"])

    for index, record in enumerate(records, 1):
        ws.append(
            [
                index,
                record.rock_type,
                record.porosity,
                record.permeability,
                record.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                username,
            ]
        )

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
